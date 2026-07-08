"""Lê uma planilha 'Controle Financ. Pessoal' (.xlsx/.xlsm) enviada e extrai os dados."""
import io

from openpyxl import load_workbook


def _d(v):
    return v.date().isoformat() if hasattr(v, "date") else None


def _s(v):
    if v is None:
        return None
    t = str(v).strip()
    return t or None


def _cell(row, i):
    return row[i] if i < len(row) else None


def parse(data: bytes) -> dict:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Não foi possível ler o arquivo. Envie um .xlsx ou .xlsm válido.")

    names = wb.sheetnames
    if "LANCAMENTO" not in names:
        raise ValueError(
            "A planilha não tem a aba 'LANCAMENTO'. Use o modelo 'Controle Financ. Pessoal'.")

    # ---- CONFIGURACAO: categorias e parâmetros ----
    categories, seen = [], set()
    receita_mensal = custo_vida = 0.0
    if "CONFIGURACAO" in names:
        ws = wb["CONFIGURACAO"]
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i <= 3:
                label = _s(_cell(row, 5))
                val = _cell(row, 6)
                if label and "Receita Mensal" in label and isinstance(val, (int, float)):
                    receita_mensal = float(val)
                if label and "Custo de Vida" in label and isinstance(val, (int, float)):
                    custo_vida = float(val)
            if i < 5:
                continue
            classificacao = _s(_cell(row, 1))
            grupo, tipo = _s(_cell(row, 2)), _s(_cell(row, 3))
            categoria, sub = _s(_cell(row, 4)), _s(_cell(row, 5))
            meta = _cell(row, 6)
            if not classificacao or not categoria or " - " not in classificacao:
                continue
            if tipo not in ("RECEITA", "DESPESA"):
                continue
            if classificacao in seen:
                continue
            seen.add(classificacao)
            categories.append({
                "classificacao": classificacao,
                "grupo": grupo or "OPERACIONAL",
                "tipo": tipo,
                "categoria": categoria,
                "subcategoria": sub,
                "meta_mes": float(meta) if isinstance(meta, (int, float)) else 0.0,
            })

    # ---- LANCAMENTO ----
    transactions = []
    ws = wb["LANCAMENTO"]
    for row in ws.iter_rows(min_row=4, values_only=True):
        desc, valor = _s(_cell(row, 3)), _cell(row, 4)
        if not desc or not isinstance(valor, (int, float)):
            continue
        transactions.append({
            "dt_compra": _d(_cell(row, 1)),
            "dt_venc": _d(_cell(row, 2)) or _d(_cell(row, 1)),
            "classificacao": desc,
            "valor": round(float(valor), 2),
            "instituicao": _s(_cell(row, 5)),
            "pessoa": _s(_cell(row, 6)),
            "status": _s(_cell(row, 7)) or "Previsto",
            "obs": _s(_cell(row, 8)),
            "grupo": _s(_cell(row, 9)),
            "tipo": _s(_cell(row, 10)),
            "categoria": _s(_cell(row, 11)),
            "subcategoria": _s(_cell(row, 12)),
        })

    # ---- PROJETOS ----
    projects = []
    fator = 6.0
    if "PROJETOS" in names:
        ws = wb["PROJETOS"]
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i == 1 and isinstance(_cell(row, 8), (int, float)):
                fator = float(_cell(row, 8))
            if i >= 2 and _s(_cell(row, 1)) and isinstance(_cell(row, 2), (int, float)):
                projects.append({
                    "descricao": _s(_cell(row, 1)),
                    "valor": float(_cell(row, 2)),
                    "ano": int(_cell(row, 3)) if isinstance(_cell(row, 3), (int, float)) else None,
                    "prazo": _s(_cell(row, 4)),
                })

    # ---- PATRIMONIO ----
    patrimonio = {"investimentos": [], "bens": [], "dividas": []}
    if "PATRIMONIO" in names:
        ws = wb["PATRIMONIO"]
        for row in ws.iter_rows(min_row=3, values_only=True):
            if _s(_cell(row, 1)) and isinstance(_cell(row, 5), (int, float)):
                patrimonio["investimentos"].append({
                    "instituicao": _s(_cell(row, 1)), "fixa_var": _s(_cell(row, 2)),
                    "prazo_projeto": _s(_cell(row, 3)), "ativo": _s(_cell(row, 4)),
                    "valor": float(_cell(row, 5)),
                })
            if _s(_cell(row, 7)) and isinstance(_cell(row, 8), (int, float)):
                patrimonio["bens"].append({
                    "descricao": _s(_cell(row, 7)), "valor": float(_cell(row, 8)),
                    "saldo_devedor": float(_cell(row, 9)) if isinstance(_cell(row, 9), (int, float)) else 0.0,
                })
            if _s(_cell(row, 12)) and any(
                    isinstance(_cell(row, j), (int, float)) for j in (13, 14, 15)):
                patrimonio["dividas"].append({
                    "descricao": _s(_cell(row, 12)),
                    "num_parcelas": int(_cell(row, 13)) if isinstance(_cell(row, 13), (int, float)) else 0,
                    "valor_parcela": float(_cell(row, 14)) if isinstance(_cell(row, 14), (int, float)) else 0.0,
                    "saldo_devedor": float(_cell(row, 15)) if isinstance(_cell(row, 15), (int, float)) else 0.0,
                })

    wb.close()
    return {
        "settings": {
            "receita_mensal": receita_mensal,
            "custo_vida_mensal": custo_vida,
            "fator_reserva": fator,
        },
        "categories": categories,
        "transactions": transactions,
        "projects": projects,
        "patrimonio": patrimonio,
    }
