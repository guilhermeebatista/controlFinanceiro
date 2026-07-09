"""Extrai os dados da planilha Controle Financ. Pessoal 6.0 para app/seed.json."""
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSM = ROOT / "Controle Financ. Pessoal 6.0 - R00.xlsm"
OUT = ROOT / "app" / "seed.json"

wb = load_workbook(XLSM, read_only=True, data_only=True)

def d(v):
    return v.date().isoformat() if hasattr(v, "date") else None

def s(v):
    return str(v).strip() if v is not None and str(v).strip() else None

# --- CONFIGURACAO: classificações ---
categories = []
seen = set()
ws = wb["CONFIGURACAO"]
for row in ws.iter_rows(min_row=5, values_only=True):
    classificacao, grupo, tipo, categoria, sub = row[1], row[2], row[3], row[4], row[5]
    meta = row[6] if len(row) > 6 else None
    if not classificacao or not categoria:
        continue
    key = s(classificacao)
    if key in seen:
        continue
    seen.add(key)
    categories.append({
        "classificacao": key,
        "grupo": s(grupo) or "OPERACIONAL",
        "tipo": s(tipo) or "DESPESA",
        "categoria": s(categoria),
        "subcategoria": s(sub),
        "meta_mes": float(meta) if isinstance(meta, (int, float)) else 0.0,
    })

# --- LANCAMENTO ---
transactions = []
ws = wb["LANCAMENTO"]
for row in ws.iter_rows(min_row=4, values_only=True):
    dt_compra, dt_venc, desc, valor = row[1], row[2], row[3], row[4]
    inst, pessoa, status, obs = row[5], row[6], row[7], row[8]
    grupo, tipo, cat, sub = row[9], row[10], row[11], row[12]
    if not desc or not isinstance(valor, (int, float)):
        continue
    transactions.append({
        "dt_compra": d(dt_compra),
        "dt_venc": d(dt_venc),
        "classificacao": s(desc),
        "valor": round(float(valor), 2),
        "instituicao": s(inst),
        "pessoa": s(pessoa),
        "status": s(status) or "Previsto",
        "obs": s(obs),
        "grupo": s(grupo) or "OPERACIONAL",
        "tipo": s(tipo) or "DESPESA",
        "categoria": s(cat),
        "subcategoria": s(sub),
    })

# --- PROJETOS ---
projects = []
ws = wb["PROJETOS"]
fator_reserva = 6
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if i == 1 and len(row) > 8 and isinstance(row[8], (int, float)):
        fator_reserva = float(row[8])
    if i >= 2 and row[1] and isinstance(row[2], (int, float)):
        projects.append({
            "descricao": s(row[1]),
            "valor": float(row[2]),
            "ano": int(row[3]) if isinstance(row[3], (int, float)) else None,
            "prazo": s(row[4]),
        })

# --- PATRIMONIO (vazio na planilha, mantém estrutura) ---
patrimonio = {"investimentos": [], "bens": [], "dividas": []}
ws = wb["PATRIMONIO"]
for row in ws.iter_rows(min_row=3, values_only=True):
    if row[1] and isinstance(row[5], (int, float)):
        patrimonio["investimentos"].append({
            "instituicao": s(row[1]), "fixa_var": s(row[2]),
            "prazo_projeto": s(row[3]), "ativo": s(row[4]), "valor": float(row[5]),
        })
    if len(row) > 8 and row[7] and isinstance(row[8], (int, float)):
        patrimonio["bens"].append({
            "descricao": s(row[7]), "valor": float(row[8]),
            "saldo_devedor": float(row[9]) if isinstance(row[9], (int, float)) else 0.0,
        })
    if len(row) > 13 and row[12]:
        patrimonio["dividas"].append({
            "descricao": s(row[12]),
            "num_parcelas": int(row[13]) if isinstance(row[13], (int, float)) else 0,
            "valor_parcela": float(row[14]) if isinstance(row[14], (int, float)) else 0.0,
            "saldo_devedor": float(row[15]) if len(row) > 15 and isinstance(row[15], (int, float)) else 0.0,
        })

# --- CONFIG geral ---
ws = wb["CONFIGURACAO"]
receita_mensal, custo_vida = 0.0, 0.0
for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
    label = s(row[5]) if len(row) > 5 else None
    val = row[6] if len(row) > 6 else None
    if label and "Receita Mensal" in label and isinstance(val, (int, float)):
        receita_mensal = float(val)
    if label and "Custo de Vida" in label and isinstance(val, (int, float)):
        custo_vida = float(val)

seed = {
    "settings": {
        "receita_mensal": receita_mensal,
        "custo_vida_mensal": custo_vida,
        "fator_reserva": fator_reserva,
    },
    "categories": categories,
    "transactions": transactions,
    "projects": projects,
    "patrimonio": patrimonio,
}

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(seed, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"categories={len(categories)} transactions={len(transactions)} projects={len(projects)}")
