"""Exporta os dados de um usuário para uma planilha re-importável via /api/import.

Gera um .xlsx no layout do modelo 'Controle Financ. Pessoal' (mesmas abas/colunas
que app/importer.py lê), para servir de backup. Rode dentro do container:

    docker cp scripts/export_backup.py minhas-contas:/tmp/export_backup.py
    docker exec minhas-contas python /tmp/export_backup.py <email> /tmp/backup.xlsx
"""
import sqlite3
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font

DB_PATH = "/data/financas.db"


def dt(v):
    if not v:
        return None
    try:
        return datetime.strptime(str(v)[:10], "%Y-%m-%d")
    except ValueError:
        return None


def main(email, out_path):
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    u = c.execute(
        "SELECT * FROM users WHERE lower(email)=? OR lower(usuario)=?",
        (email.lower(), email.lower()),
    ).fetchone()
    if not u:
        raise SystemExit(f"Usuário não encontrado: {email}")
    uid = u["id"]

    settings = {r["key"]: r["value"] for r in
                c.execute("SELECT key,value FROM settings WHERE user_id=?", (uid,))}
    cats = c.execute(
        "SELECT * FROM categories WHERE user_id=? ORDER BY tipo, categoria, subcategoria",
        (uid,)).fetchall()
    txs = c.execute(
        "SELECT * FROM transactions WHERE user_id=? ORDER BY dt_venc, id", (uid,)).fetchall()
    projs = c.execute("SELECT * FROM projects WHERE user_id=? ORDER BY ano", (uid,)).fetchall()
    invs = c.execute("SELECT * FROM investments WHERE user_id=?", (uid,)).fetchall()
    assets = c.execute("SELECT * FROM assets WHERE user_id=?", (uid,)).fetchall()
    debts = c.execute("SELECT * FROM debts WHERE user_id=?", (uid,)).fetchall()
    insts = c.execute("SELECT * FROM institutions WHERE user_id=? ORDER BY nome", (uid,)).fetchall()
    people = c.execute("SELECT * FROM people WHERE user_id=? ORDER BY nome", (uid,)).fetchall()
    c.close()

    wb = Workbook()
    bold = Font(bold=True)

    # ---------- CONFIGURACAO ----------
    # importer: labels em F1:G3 (idx5/6); dados a partir da linha 5
    # (idx1=classificacao,2=grupo,3=tipo,4=categoria,5=subcategoria,6=meta_mes)
    ws = wb.active
    ws.title = "CONFIGURACAO"
    ws["F1"] = "Receita Mensal"
    ws["G1"] = settings.get("receita_mensal", 0)
    ws["F2"] = "Custo de Vida Mensal"
    ws["G2"] = settings.get("custo_vida_mensal", 0)
    for col, h in zip("BCDEFG",
                      ["Classificacao", "Grupo", "Tipo", "Categoria", "Subcategoria", "Meta Mês"]):
        ws[f"{col}4"] = h
        ws[f"{col}4"].font = bold
    r = 5
    for cat in cats:
        ws.cell(r, 2, cat["classificacao"])
        ws.cell(r, 3, cat["grupo"])
        ws.cell(r, 4, cat["tipo"])
        ws.cell(r, 5, cat["categoria"])
        ws.cell(r, 6, cat["subcategoria"])
        ws.cell(r, 7, cat["meta_mes"] or 0)
        r += 1

    # ---------- LANCAMENTO ----------
    # importer: dados a partir da linha 4; idx1=dt_compra,2=dt_venc,3=classificacao,
    # 4=valor,5=instituicao,6=pessoa,7=status,8=obs,9=grupo,10=tipo,11=categoria,12=subcategoria
    ws = wb.create_sheet("LANCAMENTO")
    headers = ["Dt Compra", "Dt Venc", "Classificacao", "Valor", "Instituicao",
               "Pessoa", "Status", "Obs", "Grupo", "Tipo", "Categoria", "Subcategoria"]
    for i, h in enumerate(headers, start=2):
        ws.cell(3, i, h).font = bold
    r = 4
    for t in txs:
        dc = dt(t["dt_compra"])
        dv = dt(t["dt_venc"])
        if dc:
            ws.cell(r, 2, dc).number_format = "DD/MM/YYYY"
        if dv:
            ws.cell(r, 3, dv).number_format = "DD/MM/YYYY"
        ws.cell(r, 4, t["classificacao"])
        ws.cell(r, 5, round(t["valor"], 2))
        ws.cell(r, 6, t["instituicao"])
        ws.cell(r, 7, t["pessoa"])
        ws.cell(r, 8, t["status"])
        ws.cell(r, 9, t["obs"])
        ws.cell(r, 10, t["grupo"])
        ws.cell(r, 11, t["tipo"])
        ws.cell(r, 12, t["categoria"])
        ws.cell(r, 13, t["subcategoria"])
        r += 1

    # ---------- PROJETOS ----------
    # importer: fator_reserva em I1 (idx8); dados a partir da linha 2
    # (idx1=descricao,2=valor,3=ano,4=prazo)
    ws = wb.create_sheet("PROJETOS")
    ws["I1"] = settings.get("fator_reserva", 6)
    for col, h in zip("BCDE", ["Descricao", "Valor", "Ano", "Prazo"]):
        ws[f"{col}1"] = h
        ws[f"{col}1"].font = bold
    r = 2
    for p in projs:
        ws.cell(r, 2, p["descricao"])
        ws.cell(r, 3, p["valor"])
        ws.cell(r, 4, p["ano"])
        ws.cell(r, 5, p["prazo"])
        r += 1

    # ---------- PATRIMONIO ----------
    # importer: dados a partir da linha 3
    # investimentos idx1..5; bens idx7..9; dividas idx12..15
    ws = wb.create_sheet("PATRIMONIO")
    for col, h in zip("BCDEF", ["Instituicao", "Fixa/Var", "Prazo/Projeto", "Ativo", "Valor"]):
        ws[f"{col}2"] = h
        ws[f"{col}2"].font = bold
    for col, h in zip("HIJ", ["Bem", "Valor", "Saldo Devedor"]):
        ws[f"{col}2"] = h
        ws[f"{col}2"].font = bold
    for col, h in zip("MNOP", ["Dívida", "Nº Parcelas", "Valor Parcela", "Saldo Devedor"]):
        ws[f"{col}2"] = h
        ws[f"{col}2"].font = bold
    rr = 3
    for inv in invs:
        ws.cell(rr, 2, inv["instituicao"])
        ws.cell(rr, 3, inv["fixa_var"])
        ws.cell(rr, 4, inv["prazo_projeto"])
        ws.cell(rr, 5, inv["ativo"])
        ws.cell(rr, 6, inv["valor"])
        rr += 1
    rr = 3
    for a in assets:
        ws.cell(rr, 8, a["descricao"])
        ws.cell(rr, 9, a["valor"])
        ws.cell(rr, 10, a["saldo_devedor"])
        rr += 1
    rr = 3
    for d in debts:
        ws.cell(rr, 13, d["descricao"])
        ws.cell(rr, 14, d["num_parcelas"])
        ws.cell(rr, 15, d["valor_parcela"])
        ws.cell(rr, 16, d["saldo_devedor"])
        rr += 1

    # ---------- CADASTROS (referência; não é reimportada diretamente) ----------
    ws = wb.create_sheet("CADASTROS")
    ws["A1"] = "Instituições (nome / tipo / saldo inicial / descrição)"
    ws["A1"].font = bold
    r = 2
    for i in insts:
        ws.cell(r, 1, i["nome"])
        ws.cell(r, 2, i["tipo"])
        ws.cell(r, 3, i["saldo_inicial"])
        ws.cell(r, 4, i["descricao"])
        r += 1
    r += 2
    ws.cell(r, 1, "Pessoas").font = bold
    r += 1
    for p in people:
        ws.cell(r, 1, p["nome"])
        r += 1

    wb.save(out_path)
    print(f"OK -> {out_path} | lançamentos={len(txs)} categorias={len(cats)} "
          f"projetos={len(projs)} instituições={len(insts)} pessoas={len(people)}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
